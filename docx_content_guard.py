#!/usr/bin/env python3
"""比较两个 DOCX 版本并生成本地 HTML 内容差异报告。"""

import argparse
import os
import re
import html as html_mod
import unicodedata
import difflib
import webbrowser
from pathlib import Path
from docx import Document


# ============================================================
#  文本归一化
# ============================================================

def norm_ref(text):
    """参考文献编号归一化：[1] → [?]"""
    return re.sub(r'^\[\d+\]', '[?]', text)


def normalize_text(text):
    """全文归一化：消除纯格式差异（全角/半角、空白、不可见字符、标点）。"""
    s = unicodedata.normalize('NFKC', text)
    s = s.replace('​', '').replace('\xa0', ' ').replace('﻿', '')
    s = re.sub(r'\s+', ' ', s).strip()
    s = s.replace('——', '—').replace('……', '…')
    s = s.replace('“', '"').replace('”', '"')
    s = s.replace('‘', "'").replace('’', "'")
    return s


# ============================================================
#  字符级 diff
# ============================================================

def char_diff_html(original, modified):
    """对两段文本做字符级比对，返回带 <del>/<ins> 的 HTML 片段。"""
    sm = difflib.SequenceMatcher(None, original, modified)
    parts = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            parts.append(html_mod.escape(original[i1:i2]))
        elif tag == 'delete':
            parts.append(f'<del>{html_mod.escape(original[i1:i2])}</del>')
        elif tag == 'insert':
            parts.append(f'<ins>{html_mod.escape(modified[j1:j2])}</ins>')
        elif tag == 'replace':
            parts.append(f'<del>{html_mod.escape(original[i1:i2])}</del>')
            parts.append(f'<ins>{html_mod.escape(modified[j1:j2])}</ins>')
    return ''.join(parts)


def severity_label(original, modified):
    """根据改动字符占比返回严重度标签。"""
    total = max(len(original), len(modified))
    if total == 0:
        return 'minor'
    changed = 0
    sm = difflib.SequenceMatcher(None, original, modified)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != 'equal':
            changed += max(i2 - i1, j2 - j1)
    ratio = changed / total
    if ratio < 0.2:
        return 'minor'
    elif ratio < 0.8:
        return 'moderate'
    return 'major'


SEVERITY_CN = {'minor': '小修改', 'moderate': '较大改动', 'major': '整段重写'}


# ============================================================
#  文件查找
# ============================================================

def find_docx_files(folder):
    before_file = None
    after_file = None
    for f in os.listdir(folder):
        if '修改前' in f and f.endswith('.docx') and not f.startswith('~$'):
            before_file = os.path.join(folder, f)
        if '修改后' in f and f.endswith('.docx') and not f.startswith('~$'):
            after_file = os.path.join(folder, f)
    return before_file, after_file


# ============================================================
#  段落提取（含章节上下文）
# ============================================================

def extract_paragraphs(doc):
    """提取段落文本及所在章节标题。"""
    items = []
    current_section = ''
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style_name = p.style.name if p.style else ''

        is_heading = ('Heading' in style_name or 'heading' in style_name
                      or 'TOC' in style_name)

        if is_heading and text:
            current_section = text
            items.append({'index': i, 'text': '', 'section': '', 'is_heading': True})
        else:
            items.append({
                'index': i,
                'text': text,
                'section': current_section,
                'is_heading': False,
            })

    return items


# ============================================================
#  核心比对
# ============================================================

def run_comparison(before_file, after_file):
    doc_before = Document(before_file)
    doc_after = Document(after_file)

    before_items = extract_paragraphs(doc_before)
    after_items = extract_paragraphs(doc_after)

    # 只比对正文段落（排除标题行）
    before_texts = [it for it in before_items if not it['is_heading']]
    after_texts = [it for it in after_items if not it['is_heading']]

    doi_removed = 0
    ref_renumbered = 0
    normalized_count = 0
    unmatched_before = []
    unmatched_after = []

    for b_item in before_texts:
        b_text = b_item['text']
        if not b_text:
            continue

        b_full = normalize_text(b_text)
        b_ref = norm_ref(b_full)
        b_nodoi = re.sub(r'DOI：\S+\.', '', b_full)

        found = False
        for a_item in after_texts:
            a_text = a_item['text']
            if not a_text:
                continue

            a_full = normalize_text(a_text)
            a_ref = norm_ref(a_full)
            a_nodoi = re.sub(r'DOI：\S+\.', '', a_full)

            # 优先级：归一化相等 > 参考文献编号 > DOI 删除
            if b_full == a_full:
                if b_text != a_text:
                    normalized_count += 1
                found = True
                break
            if b_ref == a_ref:
                ref_renumbered += 1
                found = True
                break
            if b_nodoi == a_nodoi:
                doi_removed += 1
                found = True
                break

        if not found:
            unmatched_before.append(b_item)

    # 反过来找修改后新增的段落
    for a_item in after_texts:
        a_text = a_item['text']
        if not a_text:
            continue

        a_full = normalize_text(a_text)
        found = False
        for b_item in before_texts:
            b_text = b_item['text']
            if not b_text:
                continue
            if normalize_text(b_text) == a_full:
                found = True
                break
            if norm_ref(normalize_text(b_text)) == norm_ref(a_full):
                found = True
                break
            if re.sub(r'DOI：\S+\.', '', normalize_text(b_text)) == re.sub(r'DOI：\S+\.', '', a_full):
                found = True
                break
        if not found:
            unmatched_after.append(a_item)

    # ---- 段落配对：尝试把「消失」和「新增」的相似段落配对 ----
    paired = []
    unpaired_before = []
    remaining_after = list(unmatched_after)

    for b_item in unmatched_before:
        best_match = None
        best_score = 0.0
        for a_item in remaining_after:
            score = difflib.SequenceMatcher(None, b_item['text'], a_item['text']).ratio()
            if score > 0.5 and score > best_score:
                best_score = score
                best_match = a_item
        if best_match:
            remaining_after.remove(best_match)
            paired.append((b_item, best_match, best_score))
        else:
            unpaired_before.append(b_item)

    # ---- 统计 ----
    before_blanks = sum(1 for it in before_items if it['text'] == '' and not it['is_heading'])
    after_blanks = sum(1 for it in after_items if it['text'] == '' and not it['is_heading'])

    # ---- 组装异常项 ----
    unmatched_items = []

    for b_item, a_item, score in paired:
        severity = severity_label(b_item['text'], a_item['text'])
        diff = char_diff_html(b_item['text'], a_item['text'])
        unmatched_items.append({
            'para_index': b_item['index'] + 1,
            'section_title': b_item['section'],
            'direction': 'changed',
            'severity': severity,
            'severity_cn': SEVERITY_CN[severity],
            'original_text': b_item['text'],
            'modified_text': a_item['text'],
            'diff_html': diff,
        })

    for b_item in unpaired_before:
        unmatched_items.append({
            'para_index': b_item['index'] + 1,
            'section_title': b_item['section'],
            'direction': 'removed',
            'severity': 'major',
            'severity_cn': '整段消失',
            'original_text': b_item['text'],
            'modified_text': '',
            'diff_html': '',
        })

    for a_item in remaining_after:
        unmatched_items.append({
            'para_index': a_item['index'] + 1,
            'section_title': a_item['section'],
            'direction': 'added',
            'severity': 'major',
            'severity_cn': '整段新增',
            'original_text': '',
            'modified_text': a_item['text'],
            'diff_html': '',
        })

    return {
        'doi_removed': doi_removed,
        'ref_renumbered': ref_renumbered,
        'blank_diff': before_blanks - after_blanks,
        'normalized_count': normalized_count,
        'total_paragraphs': len(before_texts),
        'unmatched': unmatched_items,
    }


# ============================================================
#  HTML 报告
# ============================================================

def build_html(before_name, after_name, result):
    unmatched = result['unmatched']
    has_unexpected = len(unmatched) > 0
    safe_before_name = html_mod.escape(str(before_name))
    safe_after_name = html_mod.escape(str(after_name))

    verdict_color = '#EF4444' if has_unexpected else '#10B981'
    verdict_icon = '&#10060;' if has_unexpected else '&#10004;'

    if has_unexpected:
        verdict_title = f'发现 {len(unmatched)} 处需要复查'
        verdict_sub = (
            '改动中包含非格式性的内容变化，请人工逐条确认是否符合预期。'
            f'（共比对 {result["total_paragraphs"]} 段正文）'
        )
    else:
        verdict_title = '当前覆盖范围内未发现内容变化'
        verdict_sub = (
            '除已归一化的格式差异外，正文段落未发现内容层面的变化。<br>'
            '仍应人工复核表格、文本框、脚注、页眉和页脚。'
            f'（共比对 {result["total_paragraphs"]} 段正文）'
        )

    # ---- 异常列表 HTML ----
    unmatched_html = ''
    if has_unexpected:
        items_html = ''
        for item in unmatched:
            if item['direction'] == 'changed':
                direction_label = '内容改动'
            elif item['direction'] == 'removed':
                direction_label = '修改前消失'
            else:
                direction_label = '修改后新增'

            section_info = (
                f' · {html_mod.escape(item["section_title"])}'
                if item['section_title']
                else ''
            )
            meta = f'第 {item["para_index"]} 段{section_info} · {direction_label} · {item["severity_cn"]}'

            if item['diff_html']:
                body = f'<div class="diff-text">{item["diff_html"]}</div>'
            elif item['direction'] == 'removed':
                body = f'<div class="plain-text">{html_mod.escape(item["original_text"][:300])}</div>'
            else:
                body = f'<div class="plain-text">{html_mod.escape(item["modified_text"][:300])}</div>'

            items_html += f'''
                <div class="unmatched-item">
                    <div class="meta">{meta}</div>
                    {body}
                </div>'''

        unmatched_html = f'''
        <div class="section">
            <div class="section-title warning">&#9888; 发现意外变化（{len(unmatched)} 处）</div>
            <div class="unmatched-list">{items_html}
            </div>
        </div>'''

    # ---- 统计卡 ----
    stat_cards = f'''
        <div class="stat-card">
            <div class="stat-num">{result['doi_removed']}</div>
            <div class="stat-label">DOI 删除</div>
        </div>
        <div class="stat-card">
            <div class="stat-num">{result['ref_renumbered']}</div>
            <div class="stat-label">编号调整</div>
        </div>
        <div class="stat-card">
            <div class="stat-num">{result['blank_diff']}</div>
            <div class="stat-label">空段移除</div>
        </div>
        <div class="stat-card">
            <div class="stat-num">{result['normalized_count']}</div>
            <div class="stat-label">格式归一化</div>
        </div>'''

    html = f'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>大白核查 · 论文体检报告</title>
<style>
    :root {{
        --bg: #FAF7F2;
        --card-bg: #FFFFFF;
        --sub-bg: #FBF9F4;
        --border: #F3F0EA;
        --text-1: #1F2937;
        --text-2: #6B7280;
        --text-3: #9CA3AF;
        --brand: #E63946;
        --safe: #10B981;
        --warning: #EF4444;
        --warning-bg: #FFF5F5;
        --shadow: 0 1px 2px rgba(0,0,0,0.03), 0 4px 12px rgba(0,0,0,0.04);
        --radius-lg: 16px;
        --radius-md: 10px;
    }}

    * {{ box-sizing: border-box; margin: 0; padding: 0; }}

    body {{
        background: var(--bg);
        font-family: "PingFang SC", "Hiragino Sans GB", "Microsoft YaHei",
                     -apple-system, BlinkMacSystemFont, sans-serif;
        color: var(--text-1);
        -webkit-font-smoothing: antialiased;
        line-height: 1.6;
        padding: 40px 20px;
        min-height: 100vh;
    }}

    .container {{ max-width: 760px; margin: 0 auto; }}

    /* Header */
    .header {{ text-align: center; margin-bottom: 32px; }}

    .logo-row {{
        display: inline-flex; align-items: center; gap: 12px;
        margin-bottom: 6px;
    }}

    .logo-mark {{
        width: 40px; height: 40px;
        background: #FFFFFF;
        border: 2px solid var(--text-1);
        border-radius: 50%;
        position: relative;
        flex-shrink: 0;
    }}
    .logo-mark::before, .logo-mark::after {{
        content: ''; position: absolute; top: 50%;
        width: 7px; height: 7px;
        background: var(--text-1);
        border-radius: 50%;
        transform: translateY(-50%);
    }}
    .logo-mark::before {{ left: 7px; }}
    .logo-mark::after  {{ right: 7px; }}
    .logo-mark .line {{
        position: absolute; top: 50%; left: 14px; right: 14px;
        height: 2px; background: var(--text-1); transform: translateY(-50%);
    }}

    .brand-name {{
        font-size: 22px; font-weight: 600; letter-spacing: 0.5px;
    }}
    .subtitle {{
        font-size: 13px; color: var(--text-3); letter-spacing: 4px;
    }}

    /* Verdict Hero */
    .verdict {{
        background: var(--card-bg);
        border: 1px solid var(--border);
        border-radius: var(--radius-lg);
        padding: 56px 32px 48px;
        text-align: center;
        box-shadow: var(--shadow);
        margin-bottom: 16px;
    }}
    .verdict-icon {{
        width: 84px; height: 84px;
        border-radius: 50%;
        display: inline-flex; align-items: center; justify-content: center;
        font-size: 40px; color: #FFFFFF; font-weight: 700;
        margin-bottom: 24px;
        background: {verdict_color};
        box-shadow: 0 8px 24px rgba(0,0,0,0.08);
    }}
    .verdict-title {{
        font-size: 26px; font-weight: 600;
        color: {verdict_color};
        margin-bottom: 12px;
        letter-spacing: 0.5px;
    }}
    .verdict-sub {{
        font-size: 15px; color: var(--text-2);
        max-width: 440px; margin: 0 auto;
        line-height: 1.8;
    }}

    /* Section */
    .section {{
        background: var(--card-bg);
        border: 1px solid var(--border);
        border-radius: var(--radius-lg);
        padding: 28px;
        margin-bottom: 16px;
        box-shadow: var(--shadow);
    }}
    .section-title {{
        font-size: 12px; font-weight: 600;
        color: var(--text-3);
        letter-spacing: 3px;
        margin-bottom: 18px;
        text-transform: uppercase;
    }}
    .section-title.warning {{ color: var(--warning); }}

    /* 文件信息 */
    .file-row {{
        display: flex; align-items: center; gap: 20px;
        padding: 14px 0;
        border-bottom: 1px dashed var(--border);
    }}
    .file-row:last-child {{ border-bottom: none; }}
    .file-row:first-child {{ padding-top: 4px; }}

    .file-label {{
        font-size: 12px; color: var(--text-3);
        width: 50px; flex-shrink: 0; letter-spacing: 1px;
    }}
    .file-name {{
        font-size: 14px; color: var(--text-1);
        font-family: "SF Mono", "Menlo", "Consolas", monospace;
        word-break: break-all;
    }}

    /* 统计卡 */
    .stats {{
        display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px;
    }}
    .stat-card {{
        background: var(--sub-bg);
        border: 1px solid var(--border);
        border-radius: var(--radius-md);
        padding: 24px 12px;
        text-align: center;
    }}
    .stat-num {{
        font-size: 36px; font-weight: 700; color: var(--text-1);
        line-height: 1; margin-bottom: 10px;
        font-variant-numeric: tabular-nums;
    }}
    .stat-label {{
        font-size: 12px; color: var(--text-2); letter-spacing: 2px;
    }}

    /* 异常列表 */
    .unmatched-list {{ display: flex; flex-direction: column; gap: 10px; }}
    .unmatched-item {{
        background: var(--warning-bg);
        border-left: 3px solid var(--warning);
        padding: 14px 18px;
        border-radius: 0 8px 8px 0;
        font-size: 14px; color: var(--text-1);
        line-height: 1.75;
    }}
    .unmatched-item .meta {{
        font-size: 12px; color: var(--text-3);
        margin-bottom: 6px; letter-spacing: 0.5px;
    }}
    .unmatched-item .plain-text {{
        color: var(--text-2);
    }}

    /* 字符级 diff 样式 */
    .diff-text del {{
        background: #FEE2E2; color: #991B1B;
        text-decoration: line-through;
        padding: 1px 2px; border-radius: 2px;
    }}
    .diff-text ins {{
        background: #D1FAE5; color: #065F46;
        text-decoration: none;
        padding: 1px 2px; border-radius: 2px;
    }}

    /* Footer */
    .footer {{
        text-align: center; margin-top: 32px;
        font-size: 12px; color: var(--text-3); letter-spacing: 1px;
    }}

    /* 移动端 */
    @media (max-width: 600px) {{
        body {{ padding: 20px 12px; }}
        .verdict {{ padding: 40px 20px 32px; }}
        .section {{ padding: 20px; }}
        .stat-num {{ font-size: 28px; }}
        .stats {{ grid-template-columns: repeat(2, 1fr); gap: 8px; }}
    }}
</style>
</head>
<body>
<div class="container">

    <header class="header">
        <div class="logo-row">
            <div class="logo-mark"><div class="line"></div></div>
            <div class="brand-name">大白核查</div>
        </div>
        <div class="subtitle">论文体检报告</div>
    </header>

    <div class="verdict">
        <div class="verdict-icon">{verdict_icon}</div>
        <div class="verdict-title">{verdict_title}</div>
        <div class="verdict-sub">{verdict_sub}</div>
    </div>

    {unmatched_html}

    <div class="section">
        <div class="section-title">体 检 样 本</div>
        <div class="file-row">
            <div class="file-label">修改前</div>
            <div class="file-name">{safe_before_name}</div>
        </div>
        <div class="file-row">
            <div class="file-label">修改后</div>
            <div class="file-name">{safe_after_name}</div>
        </div>
    </div>

    <div class="section">
        <div class="section-title">预 期 改 动 统 计</div>
        <div class="stats">{stat_cards}
        </div>
    </div>

    <div class="footer">大白核查 · 守护每一篇论文的内容完整</div>

</div>
</body>
</html>'''
    return html


# ============================================================
#  主入口
# ============================================================

def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description='比较两个 DOCX 版本，并在本地生成 HTML 内容差异报告。',
    )
    parser.add_argument('before', nargs='?', help='修改前的 DOCX 文件')
    parser.add_argument('after', nargs='?', help='修改后的 DOCX 文件')
    parser.add_argument('-o', '--output', help='HTML 报告路径')
    parser.add_argument('--no-open', action='store_true', help='生成后不打开浏览器')
    parser.add_argument(
        '--strict',
        action='store_true',
        help='发现需要复查的变化时返回非零退出码',
    )
    args = parser.parse_args(argv)
    if bool(args.before) != bool(args.after):
        parser.error('必须同时提供修改前和修改后的 DOCX 文件。')
    return args


def resolve_inputs(args):
    if args.before and args.after:
        before_file = Path(args.before).expanduser().resolve()
        after_file = Path(args.after).expanduser().resolve()
    else:
        before, after = find_docx_files(Path.cwd())
        before_file = Path(before).resolve() if before else None
        after_file = Path(after).resolve() if after else None

    if before_file is None or not before_file.is_file():
        raise FileNotFoundError('找不到修改前的 DOCX 文件。')
    if after_file is None or not after_file.is_file():
        raise FileNotFoundError('找不到修改后的 DOCX 文件。')
    if before_file.suffix.lower() != '.docx' or after_file.suffix.lower() != '.docx':
        raise ValueError('输入文件必须是 .docx 格式。')
    return before_file, after_file


def main(argv=None):
    args = parse_args(argv)
    try:
        before_file, after_file = resolve_inputs(args)
    except (FileNotFoundError, ValueError) as exc:
        print(f'错误：{exc}')
        return 2

    before_name = before_file.name
    after_name = after_file.name

    print("大白正在体检……")
    result = run_comparison(before_file, after_file)

    html = build_html(before_name, after_name, result)

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else Path.cwd() / 'docx-content-report.html'
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html, encoding='utf-8')

    if not args.no_open:
        webbrowser.open(output_path.as_uri())
        print('报告已在浏览器中打开。')
    print(f'报告路径：{output_path}')
    print()
    print("体检结果：")
    print(f"  DOI 删除:     {result['doi_removed']} 处")
    print(f"  编号调整:     {result['ref_renumbered']} 处")
    print(f"  空段移除:     {result['blank_diff']} 个")
    print(f"  格式归一化:   {result['normalized_count']} 处")
    print(f"  比对段落:     {result['total_paragraphs']} 段")
    if result['unmatched']:
        print(f"  需要复查:     {len(result['unmatched'])} 处")
    else:
        print('  需要复查:     当前覆盖范围内未发现')
    print()
    return 1 if args.strict and result['unmatched'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
